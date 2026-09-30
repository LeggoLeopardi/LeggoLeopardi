import docx
from lxml import etree

from leggo_pipeline import paths
from leggo_pipeline.translations import build_translation_tei, read_block, split_block

NS = {"tei": "http://www.tei-c.org/ns/1.0"}


def make_docx(path, cell_paragraphs):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    cell = t.rows[1].cells[1]
    cell.paragraphs[0].text = cell_paragraphs[0]
    for p in cell_paragraphs[1:]:
        cell.add_paragraph(p)
    d.save(path)


def test_read_block_from_docx_table_keeps_text_as_is(tmp_path):
    f = tmp_path / "z.docx"
    make_docx(f, ["All’Italia", "O patria mia", "", "L’infini", "J’aime cette colline,", "  Et  cette haie", "", "Le soir"])
    src = {"docx": "z.docx", "table": 1, "start": 3, "end": 7}
    assert read_block(src, tmp_path) == ["L’infini", "J’aime cette colline,", "  Et  cette haie", ""]


def test_read_block_from_txt(tmp_path):
    (tmp_path / "t.txt").write_text("Véase también: El infinito\n\nSiempre querido\nme fue\n", encoding="utf-8")
    assert read_block({"txt": "t.txt", "start": 2, "end": 4}, tmp_path) == ["Siempre querido", "me fue"]


def test_split_block_title_stanzas_and_whitespace():
    title, stanzas = split_block(["L’infini", "", "J’aime  cette colline,", "  Et cette haie ", "", "", "Je m’assieds\xa0: ma pensée"])
    assert title == ["L’infini"]
    assert stanzas == [["J’aime cette colline,", "Et cette haie"], ["Je m’assieds\xa0: ma pensée"]]  # nbsp kept


def test_split_block_without_title_and_with_two_title_lines():
    assert split_block(["Всегда был мил", "И изгородь"], title_lines=0) == ([], [["Всегда был мил", "И изгородь"]])
    assert split_block(["XII", "", "L’INFINI.", "", "Toujours chères"], title_lines=2) == (["XII", "L’INFINI."], [["Toujours chères"]])


ENTRY = {
    "id": "fr_sainte-beuve_1844", "lang": "fr", "translator": "Charles-Augustin Sainte-Beuve", "year": 1844,
    "bibl": "«Revue des deux mondes», 1844, pp. 916-940", "form": "verse", "rights": "Pubblico dominio",
    "sources": [{"label": "Wikisource", "url": "https://fr.wikisource.org/wiki/L%E2%80%99Infini_%28trad._Sainte-Beuve%29"}],
    "notes": ["Stesso testo su Wikisource (spazi tipografici diversi)."],
}
POEM = {"n": 12, "roman": "XII", "title": "L'infinito"}


def test_translation_tei_is_valid_and_complete():
    data = build_translation_tei(ENTRY, POEM, ["L’infini"], [["J’aimai toujours", "Ce haut buisson"], ["Je m’assieds\xa0:"]])
    doc = etree.fromstring(data)
    rng = etree.RelaxNG(etree.parse(str(paths.SCHEMA)))
    assert rng.validate(etree.ElementTree(doc)), rng.error_log.last_error
    assert doc.xpath("string(//tei:titleStmt/tei:author)", namespaces=NS) == "Charles-Augustin Sainte-Beuve"
    assert doc.xpath("string(//tei:text/@xml:lang)", namespaces=NS) == "fr"
    assert doc.xpath("string(//tei:div/@corresp)", namespaces=NS) == "../../base/c12.xml#c12"
    assert doc.xpath("//tei:div/tei:head/text()", namespaces=NS) == ["L’infini"]
    assert [[l.text for l in lg] for lg in doc.xpath("//tei:div/tei:lg", namespaces=NS)] == [
        ["J’aimai toujours", "Ce haut buisson"], ["Je m’assieds\xa0:"],
    ]
    assert doc.xpath("string(//tei:availability)", namespaces=NS).strip() == "Pubblico dominio"
    assert doc.xpath("//tei:sourceDesc//tei:ptr/@target", namespaces=NS) == [ENTRY["sources"][0]["url"]]
    assert doc.xpath("string(//tei:sourceDesc/tei:bibl[1]/tei:date/@when)", namespaces=NS) == "1844"


def test_prose_translation_uses_paragraphs_and_year_is_optional():
    entry = dict(ENTRY, form="prose", year=None, sources=[], notes=[])
    doc = etree.fromstring(build_translation_tei(entry, POEM, [], [["Toujours chères me furent…", "Et je m’abîme."]]))
    assert doc.xpath("//tei:div/tei:p/text()", namespaces=NS) == ["Toujours chères me furent…", "Et je m’abîme."]
    assert doc.xpath("//tei:date", namespaces=NS) == []
    assert doc.xpath("//tei:div/tei:head", namespaces=NS) == []


def test_blank_line_after_every_verse_is_not_a_stanza_break():
    lines = ["THE INFINITE.", "", "This lonely hill", "", "And this hedge", ""]
    assert split_block(lines, stanza_breaks=False) == (["THE INFINITE."], [["This lonely hill", "And this hedge"]])


def test_docx_heading_is_the_paragraph_above_the_table(tmp_path):
    from leggo_pipeline.translations import docx_heading

    d = docx.Document()
    d.add_paragraph("1876")
    d.add_paragraph("H. F. Amiel, Les Estrangères, p. 42.")
    d.add_table(rows=2, cols=2)
    d.add_paragraph("1898")
    d.add_paragraph("W. K. Johnson, «The fortnightly review».")
    d.add_table(rows=2, cols=2)
    d.save(tmp_path / "z.docx")
    assert docx_heading({"docx": "z.docx", "table": 2}, tmp_path) == "W. K. Johnson, «The fortnightly review»."
    assert docx_heading({"docx": "z.docx", "table": 1}, tmp_path) == "H. F. Amiel, Les Estrangères, p. 42."


def test_indentation_with_nbsp_is_trimmed_but_inner_nbsp_kept():
    assert split_block(["\xa0\xa0This lonely hill\xa0", "Je m’assieds\xa0: ma pensée"], title_lines=0) == (
        [], [["This lonely hill", "Je m’assieds\xa0: ma pensée"]],
    )
